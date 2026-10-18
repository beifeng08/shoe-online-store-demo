#!/usr/bin/env python3
"""导入供应商目录 /Users/damon233/Downloads/萨洛丁款式集合.xlsx（用户真实供应链数据）。

职责（决策 #16）：
1. 解析 sheet：rows 3-31 = 29 款鞋；rows 32-42 = 11 行赠品（同名行合并为 9 款）。
2. 解析 xl/drawings/drawing1.xml 把每个内嵌 PNG 按 (row,col) 锚点归到所在产品行，
   沿用锚点在 XML 中的文档出现顺序（未再排序）；该款第一张图即封面。
3. 颜色中文名 → 英文 + 近似 hex 词典。
4. 码段（女：35-40# / 男：39-44# / 暂无尺码）→ 整档 EU 并集（直读，决策 #16）。
5. cwebp 编码（>1200px 缩到 1200 宽，q80）落 public/products/<handle>/<n>.webp。
6. 输出 src/server/catalog/data/supplier.json（seed.ts 的数据源）。

前置：pip install openpyxl；本机需 cwebp（brew install webp）。
运行：python3 scripts/import-catalog/import.py
"""

import json
import re
import shutil
import struct
import subprocess
import sys
import zipfile
from pathlib import Path

import openpyxl

SRC = Path("/Users/damon233/Downloads/萨洛丁款式集合.xlsx")
REPO = Path(__file__).resolve().parents[2]
IMG_ROOT = REPO / "public" / "products"
JSON_OUT = REPO / "src" / "server" / "catalog" / "data" / "supplier.json"

# 颜色词典：中文色名 → (EN 展示名, 近似 hex，SVG 兜底用；演示近似值)
COLORS: dict[str, tuple[str, str]] = {
    "暗黑": ("Ink Black", "#1a1a1a"),
    "雪白": ("Snow White", "#f4f4f0"),
    "象牙白": ("Ivory", "#f3ede2"),
    "白绿": ("White-Green", "#eef2e4"),
    "白蓝": ("White-Blue", "#e6eef4"),
    "薄荷绿": ("Mint", "#bfe3cd"),
    "荧光绿": ("Neon Green", "#d6f542"),
    "荧光橘": ("Neon Orange", "#ff8c1a"),
    "爱马仕橙": ("Vivid Orange", "#ff6a00"),
    "橙色": ("Orange", "#ff8c42"),
    "淡粉": ("Blush", "#f6d7d2"),
    "桃粉色": ("Peach", "#f5c8bd"),
    "柔雾粉": ("Soft Pink", "#f2c6cf"),
    "玫红": ("Rose", "#e84a5f"),
    "中国红": ("Chinese Red", "#c8102e"),
    "酒红": ("Wine", "#722f37"),
    "淡黄": ("Pale Yellow", "#f5e6a3"),
    "鹅蛋黄": ("Duckling Yellow", "#f6c445"),
    "奶酪黄": ("Cheese Yellow", "#f0d06a"),
    "黄色": ("Yellow", "#ffd400"),
    "亮黄": ("Bright Yellow", "#ffe01a"),
    "草青": ("Grass", "#8fbf6a"),
    "草绿": ("Grass Green", "#7cb342"),
    "青草绿": ("Field Green", "#66a83f"),
    "军绿": ("Olive", "#6b6b3f"),
    "土黄": ("Earth", "#9a7b4f"),
    "牛油果绿": ("Avocado", "#9cb356"),
    "湖蓝": ("Lake Blue", "#6cc3d5"),
    "浅蓝": ("Light Blue", "#a8d3e8"),
    "宝蓝": ("Royal Blue", "#2a52be"),
    "蓝色": ("Blue", "#3f7fd6"),
    "银色": ("Silver", "#c8c9cd"),
    "暗夜银": ("Midnight Silver", "#9aa1ab"),
    "紫色": ("Purple", "#8a63b8"),
    "香芋紫": ("Taro", "#b7a5d8"),
    "暗夜紫": ("Midnight Purple", "#5a4a7a"),
    "巧克力棕": ("Chocolate", "#5f3f2c"),
    "米橘": ("Beige-Orange", "#e0b398"),
    "米绿": ("Beige-Green", "#d9d5b0"),
    "青色": ("Teal", "#2fb3a3"),
    "绿色": ("Green", "#3f9d4a"),
    "卡其": ("Khaki", "#b7a47c"),
    "黑色": ("Black", "#111111"),
    "多色可选": ("Multi", "#9aa1ab"),
}
FALLBACK_EN, FALLBACK_HEX = "Colorful", "#8a8a8a"

# 赠品行名 → 英文名（中性描述命名，避开卡通 IP 名称，规避商标文本风险；原中文名留档）
GIFT_NAMES: dict[str, str] = {
    "小恐龙": "Dino Pal",
    "杰尼龟": "Turtle Pal",
    "可达鸭": "Duck Pal",
    "比卡丘": "Spark Pal",
    "奥特曼小恐龙": "Ultra Dino Pal",
    "奇妙种子": "Sprout Pal",
    "功能球": "Bounce Ball",
    "发财马": "Lucky Horse",
    "握力器": "Grip Ring",
}


def color_of(raw: str) -> tuple[str, str]:
    c = raw.strip()
    if c in COLORS:
        return COLORS[c]
    return (FALLBACK_EN, FALLBACK_HEX)


def parse_segments(seg_text: str) -> tuple[list[dict], list[int], str]:
    """'女：35-40#男：39-44#' → 段列表 + EU 整档并集 + 性别标签。'暂无尺码'/'多色可选' → 空。"""
    if not seg_text or "尺码" in seg_text or "多色" in seg_text:
        return [], [], "none"
    segments: list[dict] = []
    for m in re.finditer(r"(女|男)[:：]?\s*(\d{2})-(\d{2})\s*#?", seg_text):
        try:
            lo, hi = int(m.group(2)), int(m.group(3))
        except ValueError:
            continue
        if m.group(1) == "女":
            segments.append({"kind": "women", "lo": lo, "hi": hi})
        else:
            segments.append({"kind": "men", "lo": lo, "hi": hi})
    if not segments:  # 无法解析的码段文本：留空不猜测
        return [], [], "none"
    sizes: list[int] = []
    for s in segments:
        sizes.extend(range(s["lo"], s["hi"] + 1))
    genders = {s["kind"] for s in segments}
    label = (
        "both"
        if genders == {"men", "women"}
        else ("men" if "men" in genders else "women")
    )
    return segments, sorted(set(sizes)), label


def png_size(data: bytes) -> tuple[int, int]:
    if data[:8] != b"\x89PNG\r\n\x1a\n":
        return (0, 0)
    w, h = struct.unpack(">II", data[16:24])
    return (w, h)


def slugify(s: str) -> str:
    return re.sub(r"[^a-z0-9-]+", "-", s.lower()).strip("-") or "item"


def main() -> int:
    if not SRC.exists():
        print(f"缺少源文件: {SRC}")
        return 1
    if shutil.which("cwebp") is None:
        print("需要 cwebp（brew install webp）")
        return 1

    zf = zipfile.ZipFile(SRC)
    rels = {}
    for name in zf.namelist():
        if name.endswith("drawing1.xml.rels"):
            txt = zf.read(name).decode("utf-8")
            for m in re.finditer(r'Id="(rId\d+)"[^>]*Target="([^"]+)"', txt):
                # Target 为相对 drawings 目录的 '../media/imageN.png' → 归一化为 xl/media/…
                rels[m.group(1)] = f"xl/{m.group(2).removeprefix('../')}"
            break
    # (rowIdx, colIdx, mediaName) by anchor
    drawing: str | None = None
    for name in zf.namelist():
        if name.endswith("drawing1.xml"):
            drawing = zf.read(name).decode("utf-8")
            break
    if drawing is None:
        print("未找到 xl/drawings/drawing1.xml")
        return 1
    anchors: list[tuple[int, int, str]] = []
    for am in re.finditer(
        r"<xdr:((?:oneCell|twoCell|absolute)Anchor)[^>]*>(.*?)</xdr:\1>", drawing, re.S
    ):
        block = am.group(0)
        f = re.search(
            r"<xdr:from>.*?<xdr:col>(\d+)</xdr:col>.*?<xdr:row>(\d+)</xdr:row>",
            block,
            re.S,
        )
        embed = re.search(r'r:embed="(rId\d+)"', block)
        if f and embed:
            try:
                anchors.append(
                    (int(f.group(2)), int(f.group(1)), rels.get(embed.group(1), ""))
                )
            except ValueError:
                continue
    anchors = [(r, c, m) for (r, c, m) in anchors if m]

    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb["Sheet1"]
    rows = []
    for r in range(3, ws.max_row + 1):
        code = ws.cell(row=r, column=2).value
        if code is None:
            continue
        seg_raw = ws.cell(row=r, column=3).value
        seg_text = "" if seg_raw is None else str(seg_raw).strip()
        color_raw = ws.cell(row=r, column=9).value
        color_text = "" if color_raw is None else str(color_raw)
        colors_raw = [c for c in re.split(r"[/、]", color_text) if c.strip()]
        colors = [
            {"en": color_of(c)[0], "hex": color_of(c)[1], "cn": c.strip()}
            for c in colors_raw
        ]
        segments, sizes, gender = parse_segments(seg_text)
        imgs = [m for (rr, _c, m) in anchors if rr == r - 1] or []
        if imgs:
            seen = set()
            imgs = [
                x for x in imgs if not (x in seen or seen.add(x))
            ]  # 去重（同款重复锚定）
        kind = "shoe" if segments or seg_text == "暂无尺码" else "gift"
        rows.append(
            {
                "row": r,
                "code": str(code).strip(),
                "kind": kind,
                "seg_text": seg_text,
                "segments": segments,
                "genders": gender,
                "sizes_eu": sizes,
                "colors": colors,
                "images": imgs,
            }
        )
    shoes = [x for x in rows if x["kind"] == "shoe"]
    gift_rows = [x for x in rows if x["kind"] == "gift"]

    IMG_ROOT.mkdir(parents=True, exist_ok=True)
    total_bytes = 0

    def export_images(handle: str, media_names: list[str], folder: Path) -> list[str]:
        nonlocal total_bytes
        folder.mkdir(parents=True, exist_ok=True)
        out = []
        for i, m in enumerate(media_names, start=1):
            data = zf.read(m)  # m 已是 'xl/media/imageN.png' 全路径
            w, _h = png_size(data)
            tmp = folder / f"{i}.tmp.png"
            tmp.write_bytes(data)
            cmd = ["cwebp", "-quiet", "-q", "80", "-metadata", "none"]
            if w > 1200:
                cmd += ["-resize", "1200", "0"]
            dst = folder / f"{i}.webp"
            r = subprocess.run(cmd + [str(tmp), "-o", str(dst)], capture_output=True)
            tmp.unlink(missing_ok=True)
            if r.returncode != 0 or not dst.exists():
                print(f"cwebp 失败 {m}: {r.stderr.decode()[:120]}")
                sys.exit(1)
            total_bytes += dst.stat().st_size
            out.append(f"/products/{handle}/{i}.webp")
        return out

    shoe_entries = []
    for p in shoes:
        handle = slugify(p["code"])
        folder = IMG_ROOT / handle
        paths = export_images(handle, p["images"], folder)
        shoe_entries.append(
            {
                "handle": handle,
                "code": p["code"],
                "segments": p["segments"],
                "genders": p["genders"],
                "sizes_eu": p["sizes_eu"],
                "colors": p["colors"],
                "images": paths,
            }
        )

    # 赠品：同名行合并相册；唯一名集合按首次出现排序
    gift_groups: dict[str, list[dict]] = {}
    for p in gift_rows:
        gift_groups.setdefault(p["code"], []).append(p)
    gift_entries = []
    for name, group in gift_groups.items():
        title_en = GIFT_NAMES.get(name, name)
        handle = slugify(title_en)
        media = []
        for g in group:
            media.extend(g["images"])
        folder = IMG_ROOT / handle
        paths = export_images(handle, media, folder)
        colors = group[0]["colors"] or [{"en": "Multi", "hex": "#9aa1ab", "cn": "多色"}]
        gift_entries.append(
            {
                "handle": handle,
                "title": title_en,
                "title_cn": name,
                "colors": colors,
                "images": paths,
                "rows": [g["row"] for g in group],
            }
        )

    payload = {
        "_meta": {
            "source": "萨洛丁款式集合.xlsx",
            "note": "价格/营销名/描述为演示占位（决策 #16）；尺码数字直读 EU 整档；图片 WebP 化。",
            "shoes": len(shoe_entries),
            "gifts": len(gift_entries),
            "assets_bytes": total_bytes,
        },
        "shoes": shoe_entries,
        "gifts": gift_entries,
    }
    JSON_OUT.parent.mkdir(parents=True, exist_ok=True)
    JSON_OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n")
    print(
        f"OK shoes={len(shoe_entries)} gifts={len(gift_entries)} "
        f"assets={total_bytes / 1024:.0f}KB -> {JSON_OUT}"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
